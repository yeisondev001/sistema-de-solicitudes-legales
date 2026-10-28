"""
Historial de expedientes con el formato de EXPEDIENTES: una tabla por expediente.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

from .base import parse_fecha, MESES, set_cell_w, a_bytes

FUENTE = "Courier New"
ANCHOS = (2881, 361, 5417)

FILAS = (
    ("Tribunal",                    "tribunal"),
    ("Fecha De La Audiencia",       "fecha_audiencia"),
    ("No. Exp:",                    "expediente"),
    ("Recurrente",                  "recurrente"),
    ("Recurrido",                   "recurrido"),
    ("Agrimensor",                  "agrimensor"),
    ("Objeto",                      "objeto"),
    ("Parcela:",                    "parcela"),
    ("Fallo:",                      "fallo"),
    ("Fecha De Próxima Audiencia:", "fecha_proxima"),
)
FECHAS      = {"fecha_audiencia", "fecha_proxima"}
SIN_FORMATO = {"expediente", "parcela"}   # códigos: se copian tal cual

C = WD_ALIGN_PARAGRAPH.CENTER
J = WD_ALIGN_PARAGRAPH.JUSTIFY


def _fecha(txt: str) -> str:
    f = parse_fecha(txt)
    if f:
        return f"{f.day:02d} De {MESES[f.month]} {f.year}."
    return txt or "-"


def _capitalizar(txt: str) -> str:
    """'TRIBUNAL DE TIERRAS' / 'falta de comparecencia' → 'Tribunal De Tierras' / 'Falta De Comparecencia'."""
    if txt.isupper():
        txt = txt.lower()
    return " ".join(w[:1].upper() + w[1:] for w in txt.split(" "))


def _valor(entrada: dict, clave: str) -> str:
    v = (entrada.get(clave) or "").strip()
    if clave in FECHAS:
        return _fecha(v)
    if clave in SIN_FORMATO:
        return v
    return _capitalizar(v)


def _escribir(cell, texto, bold, align):
    par = cell.paragraphs[0]
    par.alignment = align
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after  = Pt(0)
    r = par.add_run(texto)
    r.font.name = FUENTE
    r.font.size = Pt(12)
    r.bold = bold


def _tabla(doc, numero: int, entrada: dict):
    t = doc.add_table(rows=len(FILAS) + 1, cols=3)
    t.style   = "Table Grid"
    t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.gridCol_lst, ANCHOS):
        gc.w = Twips(w)
    for row in t.rows:
        for cell, w in zip(row.cells, ANCHOS):
            set_cell_w(cell, w)

    for i, (etiqueta, clave) in enumerate(FILAS, start=1):
        c1, c2, c3 = t.rows[i].cells
        _escribir(c1, etiqueta, True, C)
        _escribir(c2, ":", True, J)
        _escribir(c3, _valor(entrada, clave), False, J)

    # Fila "No.XX": una sola celda; el resto de la fila queda sin bordes
    _escribir(t.rows[0].cells[0], f"No.{numero:02d}", True, C)
    tr = t.rows[0]._tr
    for tc in tr.tc_lst[1:]:
        tr.remove(tc)
    trPr = tr.get_or_add_trPr()
    grid_after = OxmlElement("w:gridAfter"); grid_after.set(qn("w:val"), "2")
    w_after    = OxmlElement("w:wAfter")
    w_after.set(qn("w:w"), str(ANCHOS[1] + ANCHOS[2])); w_after.set(qn("w:type"), "dxa")
    trPr.append(grid_after); trPr.append(w_after)


def generar(entradas: list) -> bytes:
    doc = Document()
    s = doc.sections[0]
    s.page_width,  s.page_height   = Twips(11906), Twips(16838)
    s.top_margin,  s.bottom_margin = Twips(1417),  Twips(1417)
    s.left_margin, s.right_margin  = Twips(1701),  Twips(1701)
    doc.styles["Normal"].font.name = FUENTE
    doc.styles["Normal"].font.size = Pt(12)

    for n, entrada in enumerate(entradas, start=1):
        if n > 1:
            doc.add_paragraph(); doc.add_paragraph()
        _tabla(doc, n, entrada)
    return a_bytes(doc)
